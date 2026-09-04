import csv
import gzip
import json
import sys
from collections import Counter, defaultdict
from datetime import date, datetime, timezone
from pathlib import Path

from openpyxl import load_workbook


def clean_text(value):
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def clean_integer(value):
    if value is None or value == "":
        return None
    return int(value)


def clean_date(value):
    if isinstance(value, (datetime, date)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    text = clean_text(value)
    return text[:10] if text else None


def main():
    if len(sys.argv) != 2:
        raise SystemExit("Usage: python scripts/build_data.py path/to/source.xlsx")

    source_path = Path(sys.argv[1]).resolve()
    project_root = Path(__file__).resolve().parents[1]
    output_dir = project_root / "public" / "data"
    output_dir.mkdir(parents=True, exist_ok=True)

    workbook = load_workbook(source_path, read_only=True, data_only=True)
    sheet = workbook["parsed_rows"]
    rows = sheet.iter_rows(values_only=True)
    headers = list(next(rows))
    header_index = {name: index for index, name in enumerate(headers) if name}

    def value(row, name):
        index = header_index[name]
        return row[index] if index < len(row) else None

    records = []
    invalid_target_values = 0

    for row in rows:
        record_id = clean_integer(value(row, "count"))
        if record_id is None:
            continue

        raw_target = value(row, "target_short")
        if raw_target is not None and not isinstance(raw_target, str):
            invalid_target_values += 1
            target = None
        else:
            target = clean_text(raw_target)

        records.append({
            "id": record_id,
            "date": clean_date(value(row, "date")),
            "year": clean_integer(value(row, "year")),
            "semester": clean_text(value(row, "semester")),
            "kind": clean_text(value(row, "kind")),
            "text": clean_text(value(row, "text_full")),
            "target": target,
            "sender": clean_text(value(row, "sender_short")),
            "sourcePdf": clean_text(value(row, "pl2_pdf")),
            "newspaperPage": clean_integer(value(row, "page_in_breeze")),
            "pdfPage": clean_integer(value(row, "page_in_pdf")),
            "positionOnPage": clean_integer(value(row, "count_on_page")),
        })

    records.sort(key=lambda record: record["id"])

    record_ids = [record["id"] for record in records]
    if len(record_ids) != len(set(record_ids)):
        raise ValueError("Duplicate record IDs found in parsed_rows")
    required_fields = ("date", "year", "kind", "text")
    incomplete = [record["id"] for record in records if any(record[field] is None for field in required_fields)]
    if incomplete:
        raise ValueError(f"Records missing required web fields: {incomplete[:10]}")

    kind_counts = Counter(record["kind"] for record in records)
    years = defaultdict(Counter)
    issues = defaultdict(Counter)
    for record in records:
        years[record["year"]][record["kind"]] += 1
        issues[record["date"]][record["kind"]] += 1

    def count_row(key, counts):
        darts = counts.get("DART", 0)
        pats = counts.get("PAT", 0)
        combined = counts.get("DART AND PAT", 0)
        return {
            "key": key,
            "darts": darts,
            "pats": pats,
            "combined": combined,
            "total": darts + pats + combined,
        }

    summary = {
        "generatedAt": datetime.now(timezone.utc).replace(microsecond=0).isoformat(),
        "sourceFile": source_path.name,
        "recordCount": len(records),
        "dateMin": min(record["date"] for record in records if record["date"]),
        "dateMax": max(record["date"] for record in records if record["date"]),
        "kindCounts": dict(sorted(kind_counts.items())),
        "invalidTargetValuesRemoved": invalid_target_values,
        "years": [count_row(year, years[year]) for year in sorted(years)],
        "issues": [count_row(issue_date, issues[issue_date]) for issue_date in sorted(issues)],
    }
    if sum(summary["kindCounts"].values()) != len(records):
        raise ValueError("Kind counts do not reconcile to the record count")
    if sum(item["total"] for item in summary["years"]) != len(records):
        raise ValueError("Year counts do not reconcile to the record count")
    if sum(item["total"] for item in summary["issues"]) != len(records):
        raise ValueError("Issue counts do not reconcile to the record count")

    records_json_path = output_dir / "records.json"
    summary_json_path = output_dir / "summary.json"
    csv_gzip_path = output_dir / "records.csv.gz"

    records_json_path.write_text(
        json.dumps(records, ensure_ascii=False, separators=(",", ":")),
        encoding="utf-8",
    )
    summary_json_path.write_text(
        json.dumps(summary, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

    fieldnames = list(records[0].keys())
    with gzip.open(csv_gzip_path, "wt", encoding="utf-8", newline="", compresslevel=9) as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(records)

    print(json.dumps({
        "records": len(records),
        "dateMin": summary["dateMin"],
        "dateMax": summary["dateMax"],
        "kindCounts": summary["kindCounts"],
        "invalidTargetValuesRemoved": invalid_target_values,
        "files": {
            path.name: path.stat().st_size
            for path in (records_json_path, summary_json_path, csv_gzip_path)
        },
    }, indent=2))


if __name__ == "__main__":
    main()
